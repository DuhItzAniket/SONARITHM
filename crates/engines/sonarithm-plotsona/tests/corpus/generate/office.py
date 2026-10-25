# SPDX-FileCopyrightText: 2026 SONARITHM contributors
# SPDX-License-Identifier: MIT OR Apache-2.0
"""Generates the Office part of PLOTSONA's exact-source chart corpus (DS-007, C01-T05/T11).

Charts with known data in XLSX (XlsxWriter), PPTX (python-pptx, which embeds a workbook)
and DOCX (the same chart part and workbook placed in a python-docx document), with a
ground-truth JSON file beside each. Data: the project's example datasets DS-001 to DS-004
(CC BY 4.0, examples/data). Dev-only tools, licences checked in C01-T02:
XlsxWriter 3.3.0 (BSD-2-Clause), python-pptx 1.0.2 and python-docx 1.2.0 (MIT).

Run from the repository root:  python crates/engines/sonarithm-plotsona/tests/corpus/generate/office.py
The outputs are committed; re-run only to change the corpus, and say why in the PR.
"""

import csv
import datetime as dt
import io
import json
import re
import sys
import zipfile
from pathlib import Path

import docx
import xlsxwriter
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[6]
DATA = ROOT / "examples" / "data"
OUT = Path(__file__).resolve().parents[1] / "office"
FIXED = dt.datetime(2026, 10, 26, 0, 0, 0)
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def rows(name):
    with open(DATA / name, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def datasets():
    """Each chart: (id, chart type, title, x title, y title, x labels or numbers, series)."""
    regions = rows("ds-003-regions.csv")
    rain = {}
    for r in regions:
        rain.setdefault(r["region"], [0.0] * 12)[int(r["month"]) - 1] = float(r["rain"])
    fruit = rows("ds-004-fruit.csv")
    temps = rows("ds-001-temperature.csv")[:60]
    prices = rows("ds-002-prices.csv")
    return [
        ("line", "line", "Rain by region", "Month", "Rain (mm)", MONTHS,
         [(k.title(), v) for k, v in rain.items()]),
        ("bar", "bar", "Fruit sold", "Fruit", "Sold",
         [r["fruit"] for r in fruit], [("Sold", [float(r["sold"]) for r in fruit])]),
        ("grouped-bar", "bar", "Rain, first quarter", "Month", "Rain (mm)", MONTHS[:4],
         [(k.title(), v[:4]) for k, v in rain.items()]),
        ("stacked-bar", "stacked-bar", "Rain by month", "Month", "Rain (mm)", MONTHS,
         [(k.title(), v) for k, v in rain.items()]),
        ("scatter", "scatter", "Daily temperature", "Day", "Temperature (°C)",
         [float(r["day"]) for r in temps], [("Temperature", [float(r["temp"]) for r in temps])]),
        ("pie", "pie", "Fruit sold", None, None,
         [r["fruit"] for r in fruit], [("Sold", [float(r["sold"]) for r in fruit])]),
        ("area", "area", "Weekly price", "Week", "Price ($)",
         [r["week"] for r in prices], [("Price", [float(r["price"]) for r in prices])]),
    ]


def truth(chart_type, title, x, series, x_kind="category"):
    return {
        "chart_type": chart_type,
        "title": title,
        "provenance": "exact",
        "x_kind": x_kind,
        "series": [{"name": n, "x": list(x), "y": list(v)} for n, v in series],
    }


XLSX_TYPES = {
    "line": ({"type": "line"}),
    "bar": ({"type": "column"}),
    "grouped-bar": ({"type": "column"}),
    "stacked-bar": ({"type": "column", "subtype": "stacked"}),
    "scatter": ({"type": "scatter"}),
    "pie": ({"type": "pie"}),
    "area": ({"type": "area"}),
}


def write_xlsx(path, ident, title, x_title, y_title, x, series):
    book = xlsxwriter.Workbook(str(path))
    book.set_properties({"created": FIXED})
    sheet = book.add_worksheet("Data")
    sheet.write(0, 0, "x")
    for i, value in enumerate(x):
        sheet.write(i + 1, 0, value)
    for j, (name, values) in enumerate(series):
        sheet.write(0, j + 1, name)
        for i, value in enumerate(values):
            sheet.write(i + 1, j + 1, value)
    chart = book.add_chart(XLSX_TYPES[ident])
    n = len(x)
    for j, (name, _) in enumerate(series):
        chart.add_series({
            "name": ["Data", 0, j + 1],
            "categories": ["Data", 1, 0, n, 0],
            "values": ["Data", 1, j + 1, n, j + 1],
        })
    chart.set_title({"name": title})
    if x_title:
        chart.set_x_axis({"name": x_title})
        chart.set_y_axis({"name": y_title})
    sheet.insert_chart("F2", chart)
    book.close()


def write_date_xlsx(path):
    """A line chart over a date axis (DS-002 weeks as dates from 2026-01-05)."""
    prices = rows("ds-002-prices.csv")
    start = dt.date(2026, 1, 5)
    dates = [start + dt.timedelta(weeks=int(r["week"])) for r in prices]
    values = [float(r["price"]) for r in prices]
    book = xlsxwriter.Workbook(str(path))
    book.set_properties({"created": FIXED})
    sheet = book.add_worksheet("Prices")
    fmt = book.add_format({"num_format": "yyyy-mm-dd"})
    sheet.write(0, 0, "Week")
    sheet.write(0, 1, "Price")
    for i, (d, v) in enumerate(zip(dates, values)):
        sheet.write_datetime(i + 1, 0, dt.datetime(d.year, d.month, d.day), fmt)
        sheet.write(i + 1, 1, v)
    chart = book.add_chart({"type": "line"})
    chart.add_series({
        "name": ["Prices", 0, 1],
        "categories": ["Prices", 1, 0, len(dates), 0],
        "values": ["Prices", 1, 1, len(dates), 1],
    })
    chart.set_x_axis({"date_axis": True, "num_format": "yyyy-mm-dd"})
    chart.set_title({"name": "Weekly price"})
    sheet.insert_chart("D2", chart)
    book.close()
    t = truth("line", "Weekly price", [d.isoformat() for d in dates], [("Price", values)], "date")
    return t


PPTX_TYPES = {
    "line": XL_CHART_TYPE.LINE,
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "grouped-bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked-bar": XL_CHART_TYPE.COLUMN_STACKED,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "pie": XL_CHART_TYPE.PIE,
    "area": XL_CHART_TYPE.AREA,
}


def write_pptx(ident, title, x_title, y_title, x, series):
    deck = Presentation()
    deck.core_properties.created = FIXED
    deck.core_properties.modified = FIXED
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if ident == "scatter":
        data = XyChartData()
        for name, values in series:
            s = data.add_series(name)
            for xv, yv in zip(x, values):
                s.add_data_point(xv, yv)
    else:
        data = CategoryChartData()
        data.categories = x
        for name, values in series:
            data.add_series(name, values)
    frame = slide.shapes.add_chart(PPTX_TYPES[ident], Inches(1), Inches(1), Inches(8), Inches(5), data)
    chart = frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    if x_title and ident != "pie":
        chart.category_axis.has_title = True
        chart.category_axis.axis_title.text_frame.text = x_title
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = y_title
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


CHART_CT = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
CHART_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"
INLINE = (
    '<w:p><w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
    '<wp:extent cx="5486400" cy="3200400"/><wp:docPr id="1" name="Chart 1"/>'
    '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
    '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdChart1"/>'
    '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'
)


def write_docx(path, pptx_bytes, title):
    """Places the PPTX chart part and its embedded workbook in a Word document."""
    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as deck:
        chart_xml = deck.read("ppt/charts/chart1.xml")
        rels = deck.read("ppt/charts/_rels/chart1.xml.rels").decode("utf-8")
        target = re.search(r'Target="\.\./embeddings/([^"]+)"', rels).group(1)
        workbook = deck.read(f"ppt/embeddings/{target}")
    document = docx.Document()
    document.core_properties.created = FIXED
    document.core_properties.modified = FIXED
    document.add_paragraph(title)
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source, \
            zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            body = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                body = body.decode("utf-8").replace(
                    "</Types>",
                    f'<Override PartName="/word/charts/chart1.xml" ContentType="{CHART_CT}"/>'
                    '<Default Extension="xlsx" ContentType="application/'
                    'vnd.openxmlformats-officedocument.spreadsheetml.sheet"/></Types>',
                ).encode("utf-8")
            elif item.filename == "word/_rels/document.xml.rels":
                body = body.decode("utf-8").replace(
                    "</Relationships>",
                    f'<Relationship Id="rIdChart1" Type="{CHART_REL}" Target="charts/chart1.xml"/>'
                    "</Relationships>",
                ).encode("utf-8")
            elif item.filename == "word/document.xml":
                body = body.decode("utf-8").replace("<w:sectPr", INLINE + "<w:sectPr", 1).encode("utf-8")
            out.writestr(item, body)
        out.writestr("word/charts/chart1.xml", chart_xml)
        out.writestr("word/charts/_rels/chart1.xml.rels", rels)
        out.writestr(f"word/embeddings/{target}", workbook)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    count = 0
    for ident, chart_type, title, x_title, y_title, x, series in datasets():
        x_kind = "number" if ident == "scatter" else "category"
        t = truth(chart_type, title, x, series, x_kind)
        write_xlsx(OUT / f"xlsx-{ident}.xlsx", ident, title, x_title, y_title, x, series)
        pptx_bytes = write_pptx(ident, title, x_title, y_title, x, series)
        (OUT / f"pptx-{ident}.pptx").write_bytes(pptx_bytes)
        write_docx(OUT / f"docx-{ident}.docx", pptx_bytes, title)
        for kind in ("xlsx", "pptx", "docx"):
            (OUT / f"{kind}-{ident}.truth.json").write_text(json.dumps(t, indent=1) + "\n", encoding="utf-8", newline="\n")
            count += 1
    t = write_date_xlsx(OUT / "xlsx-date-line.xlsx")
    (OUT / "xlsx-date-line.truth.json").write_text(json.dumps(t, indent=1) + "\n", encoding="utf-8", newline="\n")
    count += 1
    print(f"{count} Office charts in {OUT.relative_to(ROOT)}", file=sys.stderr)


if __name__ == "__main__":
    main()
